import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.services.enterprise_evaluation_dataset import CASES, CATEGORY_LABELS, summarize_dataset
from app.core.loader import load_document
from app.core.splitter import split_documents
from langchain_core.documents import Document

PROJECT_ROOT = Path(__file__).resolve().parents[2]
ENTERPRISE_DOC_DIR = PROJECT_ROOT / "data" / "enterprise_scale"
MIN_PAGE_EQUIVALENT = 20


def test_enterprise_scale_dataset_has_required_fields():
    summary = summarize_dataset()

    assert summary["version"] == "enterprise_scale_v1"
    assert summary["case_count"] >= 30
    assert set(summary["categories"]) == set(CATEGORY_LABELS)
    for case in CASES:
        assert case.case_id
        assert case.category in CATEGORY_LABELS
        assert case.question
        assert case.answer_points
        if case.category != "E":
            assert case.expected_doc
            assert case.expected_section
            assert case.expected_clause


def test_enterprise_scale_dataset_covers_all_categories():
    summary = summarize_dataset()

    assert summary["categories"]["A"]["case_count"] >= 10
    assert summary["categories"]["B"]["case_count"] >= 5
    assert summary["categories"]["C"]["case_count"] >= 5
    assert summary["categories"]["D"]["case_count"] >= 5
    assert summary["categories"]["E"]["case_count"] >= 5


def test_enterprise_scale_files_use_mixed_supported_formats():
    files = sorted(path for path in ENTERPRISE_DOC_DIR.iterdir() if path.is_file())
    suffixes = {path.suffix.lower() for path in files}

    assert len(files) >= 12
    assert {".txt", ".pdf", ".docx", ".csv", ".xlsx"}.issubset(suffixes)
    assert ".doc" not in suffixes


def test_enterprise_scale_files_are_parseable_and_contain_expected_clauses():
    files_by_stem = {path.stem: path for path in ENTERPRISE_DOC_DIR.iterdir() if path.is_file()}
    text_by_doc: dict[str, str] = {}

    for doc_name, path in files_by_stem.items():
        loaded = load_document(str(path), path.name)
        joined = "\n".join(item.page_content for item in loaded)
        text_by_doc[doc_name] = joined
        assert doc_name in joined
        assert len(joined) > 1200

    for case in CASES:
        if case.category == "E":
            continue
        assert case.expected_doc in text_by_doc
        doc_text = text_by_doc[case.expected_doc]
        assert case.expected_section in doc_text
        assert case.expected_clause in doc_text


def test_enterprise_scale_files_have_at_least_twenty_page_equivalent_content():
    from docx import Document
    from openpyxl import load_workbook
    from pypdf import PdfReader

    for path in ENTERPRISE_DOC_DIR.iterdir():
        if not path.is_file():
            continue
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            assert len(PdfReader(str(path)).pages) >= MIN_PAGE_EQUIVALENT
        elif suffix == ".docx":
            doc = Document(str(path))
            joined_xml = "\n".join(paragraph._p.xml for paragraph in doc.paragraphs)
            assert joined_xml.count('w:type="page"') >= MIN_PAGE_EQUIVALENT - 1
            assert f"第{MIN_PAGE_EQUIVALENT:02d}页" in "\n".join(p.text for p in doc.paragraphs)
        elif suffix == ".xlsx":
            workbook = load_workbook(str(path), read_only=True, data_only=True)
            assert workbook.active.max_row >= 100
            page_sheets = [name for name in workbook.sheetnames if name.startswith("P")]
            assert len(page_sheets) >= MIN_PAGE_EQUIVALENT
        elif suffix in {".txt", ".csv"}:
            text = path.read_text(encoding="utf-8-sig")
            assert f"第{MIN_PAGE_EQUIVALENT:02d}页" in text
        else:
            raise AssertionError(f"unexpected file type: {path.name}")


def test_enterprise_scale_files_are_not_old_repeated_placeholder_text():
    combined_lines: list[str] = []
    for path in ENTERPRISE_DOC_DIR.iterdir():
        if not path.is_file() or path.suffix.lower() not in {".txt", ".csv"}:
            continue
        combined_lines.extend(line.strip() for line in path.read_text(encoding="utf-8-sig").splitlines() if line.strip())

    assert "本条用于描述" not in "\n".join(combined_lines)
    repeated_ratio = 1 - (len(set(combined_lines)) / max(len(combined_lines), 1))
    assert repeated_ratio < 0.35


def test_enterprise_scale_dataset_has_no_unintended_policy_conflicts():
    files_by_stem = {path.stem: path for path in ENTERPRISE_DOC_DIR.iterdir() if path.is_file()}
    text_by_doc = {
        doc_name: "\n".join(item.page_content for item in load_document(str(path), path.name))
        for doc_name, path in files_by_stem.items()
    }

    assert "工资发放日遇节假日提前至最近一个工作日" in text_by_doc["员工手册"]
    assert "工资日遇节假日提前至最近一个工作日" in text_by_doc["薪酬绩效制度"]
    assert "工资发放日遇节假日顺延" not in "\n".join(text_by_doc.values())
    assert "工资日遇节假日顺延" not in "\n".join(text_by_doc.values())

    expected_single_policy_markers = {
        "一线城市普通员工住宿上限为480元/晚": "差旅与费用报销制度",
        "预算外付款超过10万元需CFO审批": "财务付款与预算制度",
        "P1生产事故须在10分钟内首次响应": "IT运维与服务SLA",
        "S1客户故障需30分钟内给出处理方案": "客户支持与售后制度",
        "生产发布窗口默认安排在周二、周四20:00-22:00": "研发规范与发布流程",
        "访客临时门禁最长有效期为8小时": "行政后勤制度",
        "战略供应商准入至少完成资质、财务、安全、合规四项审查": "采购与供应商管理制度",
    }
    for marker, doc_name in expected_single_policy_markers.items():
        docs_with_marker = [name for name, text in text_by_doc.items() if marker in text]
        assert docs_with_marker == [doc_name]


def test_splitter_adds_structured_source_section_and_clause_header():
    docs = [
        Document(
            page_content=(
                "# 员工手册\n"
                "## 福利与关怀\n"
                "HR-05-004 年度体检安排在每年10月完成，由人力资源部统一预约。\n"
                "HR-05-008 生日福利以电子券形式发放，不折现、不跨年补发。\n"
            ),
            metadata={"source": "员工手册.docx"},
        )
    ]

    chunks = split_documents(docs)
    birthday = next(chunk.page_content for chunk in chunks if "HR-05-008" in chunk.page_content)

    assert birthday.startswith("来源：员工手册.docx > 福利与关怀 > HR-05-008")
    assert "生日福利以电子券形式发放" in birthday
    assert "HR-05-004" not in birthday


def test_splitter_prefers_page_section_over_bare_doc_title_heading():
    docs = [
        Document(
            page_content=(
                "员工手册\n"
                "第01页 员工手册 - 入职管理\n"
                "HR-01-001 新员工入职前须完成身份核验、学历核验和保密承诺签署。\n"
            ),
            metadata={"source": "员工手册.docx"},
        )
    ]

    chunks = split_documents(docs)

    assert chunks[0].page_content.startswith("来源：员工手册.docx > 入职管理 > HR-01-001")
