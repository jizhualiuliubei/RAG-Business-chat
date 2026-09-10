"""Generate mixed-format enterprise-scale RAG evaluation documents.

The corpus is synthetic but intentionally structured like real company
materials: policies, tables, approval matrices, exception notes, SLA rows and
cross-document references. Every generated file contains stable clause IDs used
by enterprise_scale_v1 evaluation cases.
"""
from __future__ import annotations

import csv
import shutil
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "enterprise_scale"
PAGE_COUNT = 20


@dataclass(frozen=True)
class Section:
    title: str
    clauses: tuple[str, ...]


@dataclass(frozen=True)
class EnterpriseDoc:
    name: str
    suffix: str
    owner: str
    scope: str
    sections: tuple[Section, ...]


DOCS: tuple[EnterpriseDoc, ...] = (
    EnterpriseDoc(
        "员工手册",
        ".docx",
        "人力资源部",
        "适用于正式员工、试用期员工和外包驻场人员的日常管理。",
        (
            Section("入职管理", ("HR-01-001 新员工入职前须完成身份核验、学历核验和保密承诺签署。", "HR-01-004 试用期导师需在第15天、第45天和转正前提交三次反馈。")),
            Section("考勤休假", ("HR-02-002 弹性到岗范围为8:30-10:00，核心协作时间为10:00-17:00。", "HR-02-006 连续请病假超过3个工作日应上传医院证明。")),
            Section("福利与关怀", ("HR-05-004 年度体检安排在每年10月完成，由人力资源部统一预约，正式员工可参加。", "HR-05-008 生日福利以电子券形式发放，不折现、不跨年补发。")),
            Section("薪酬发放", ("HR-06-002 员工手册规定工资发放日遇节假日提前至最近一个工作日发放，并以薪酬绩效制度为准。", "HR-06-005 薪酬异议应先由直属负责人确认考勤和绩效事实。")),
            Section("离职管理", ("HR-08-003 员工离职需完成工作交接、资产归还和保密确认。", "HR-08-006 离职交接清单须由接收人、直属负责人和人力资源三方确认。")),
        ),
    ),
    EnterpriseDoc(
        "信息安全管理制度",
        ".pdf",
        "信息安全委员会",
        "覆盖数据分级、账号安全、终端安全、第三方访问和安全事件响应。",
        (
            Section("数据分级与外发", ("SEC-02-001 数据分为公开、内部、秘密、机密四级，客户资料默认不低于秘密级。", "SEC-02-006 秘密级数据外发前需部门负责人审批，并由信息安全负责人复核，禁止使用个人邮箱或微信外发。")),
            Section("账号安全", ("SEC-03-002 管理员账号必须启用多因素认证，权限申请有效期最长90天。", "SEC-03-008 共享账号须登记责任人，不得用于生产系统变更。")),
            Section("终端安全", ("SEC-04-004 终端丢失应在2小时内向IT和信息安全报备。", "SEC-04-009 研发终端接入生产堡垒机前须通过基线检查。")),
            Section("第三方访问", ("SEC-06-003 供应商远程访问生产系统须通过临时VPN、最小权限和全程审计。", "SEC-06-006 第三方访问日志保留不少于180天。")),
            Section("安全事件", ("SEC-07-002 涉及客户数据泄露的事件须在30分钟内升级到信息安全负责人。", "SEC-07-006 事件复盘需包含根因、影响范围、补救措施和负责人。")),
        ),
    ),
    EnterpriseDoc(
        "IT运维与服务SLA",
        ".docx",
        "IT 运维部",
        "规定服务台、生产事故、变更窗口、备份巡检和资产管理的服务水平。",
        (
            Section("服务台支持", ("IT-01-001 服务台工作日响应时间为9:00-18:30。", "IT-01-006 影响单人的普通办公问题按P4处理。")),
            Section("生产事故分级", ("IT-03-001 P1生产事故须在10分钟内首次响应，30分钟内建立故障群并指定负责人。", "IT-03-004 P2事故须在30分钟内首次响应，2小时内给出临时绕行方案。")),
            Section("变更窗口", ("IT-04-001 生产变更默认安排在周二和周四晚间窗口。", "IT-04-002 离职账号应在离职生效后24小时内关闭或冻结。")),
            Section("备份巡检", ("IT-05-003 核心数据库每日增量备份，每周至少一次恢复演练。", "IT-05-007 备份失败连续两次需升级运维经理。")),
            Section("资产管理", ("IT-06-002 服务器资产标签、用途、负责人和到期时间须同步CMDB。", "IT-06-008 闲置云资源连续14天无访问记录应提交释放评估。")),
        ),
    ),
    EnterpriseDoc(
        "采购与供应商管理制度",
        ".xlsx",
        "采购部",
        "管理供应商准入、采购审批、验收分离、软件订阅采购和供应商评估。",
        (
            Section("供应商准入", ("PUR-02-003 战略供应商准入至少完成资质、财务、安全、合规四项审查。", "PUR-02-006 涉及客户数据的供应商须额外完成信息安全评估。")),
            Section("采购审批", ("PUR-03-001 预算内5万元以下采购需直属负责人审批。", "PUR-03-002 预算内5万至20万元采购需部门负责人和采购经理审批。", "PUR-03-004 超过20万元采购需总经理审批。")),
            Section("验收分离", ("PUR-05-003 采购验收人不得与申请人为同一人，应保持职责分离。", "PUR-05-007 验收记录至少保留合同、发票、交付清单和验收截图。")),
            Section("软件订阅采购", ("PUR-06-002 软件订阅采购须确认供应商、授权范围、续费周期和数据安全条款。", "PUR-06-006 SaaS订阅超过一年须评估退出机制和数据导出方式。")),
            Section("供应商评估", ("PUR-08-002 年度供应商评分低于70分应进入观察名单。", "PUR-08-006 重大违约供应商一年内不得参与同类采购。")),
        ),
    ),
    EnterpriseDoc(
        "差旅与费用报销制度",
        ".pdf",
        "财务共享中心",
        "覆盖交通、住宿、餐补、报销材料和特殊差旅场景。",
        (
            Section("交通标准", ("TRV-02-001 高铁二等座为普通员工默认交通标准。", "TRV-02-006 夜间抵达后打车费用可按实际票据报销。")),
            Section("住宿标准", ("TRV-03-002 一线城市普通员工住宿上限为480元/晚，经理级为650元/晚。", "TRV-03-006 超标准住宿需在出行前获得部门负责人批准。")),
            Section("餐饮补贴", ("TRV-04-001 国内差旅补贴为120元/天，市内交通上限60元/天。", "TRV-04-005 已由会议主办方提供餐饮的，不再重复领取餐补。")),
            Section("报销材料", ("TRV-05-002 报销需提交审批单、发票、行程单和付款记录。", "TRV-05-006 电子发票须完成查重校验后入账。")),
            Section("特殊差旅", ("TRV-06-003 客户紧急现场支持可先出行后补审批，但需24小时内补齐记录。", "TRV-06-009 境外差旅须额外完成保险和外事备案。")),
        ),
    ),
    EnterpriseDoc(
        "薪酬绩效制度",
        ".xlsx",
        "薪酬绩效委员会",
        "管理薪酬结构、绩效周期、绩效系数、薪酬发放和绩效申诉。",
        (
            Section("薪酬结构", ("PAY-01-001 薪酬由基本工资、岗位工资、绩效奖金和专项补贴构成。", "PAY-01-006 调薪窗口原则上安排在每年4月和10月。")),
            Section("绩效周期", ("PAY-03-002 季度绩效用于过程反馈，年度绩效用于奖金和晋升参考。", "PAY-03-006 绩效面谈需保留员工确认记录。")),
            Section("绩效系数", ("PAY-05-004 P3员工绩效为A时绩效系数为1.15。", "PAY-05-006 P4员工绩效为B时绩效系数为1.00。")),
            Section("绩效申诉", ("PAY-06-002 年度绩效申诉应在结果公示后5个工作日内提交。", "PAY-06-006 申诉复核由隔级负责人、人力资源和业务BP共同完成。")),
            Section("薪酬发放", ("PAY-07-001 薪酬绩效制度规定工资日遇节假日提前至最近一个工作日发放。", "PAY-07-004 补发、扣回和专项奖金需在工资条中单独列示。")),
        ),
    ),
    EnterpriseDoc(
        "培训与晋升制度",
        ".docx",
        "人才发展部",
        "覆盖年度培训、必修课程、外部培训、晋升评审和导师机制。",
        (
            Section("培训计划", ("TRN-01-001 年度培训计划由业务部门提出需求，人力资源汇总预算。", "TRN-01-005 新经理训练营每半年组织一次。")),
            Section("必修课程", ("TRN-02-003 信息安全、合规和反舞弊课程为全员年度必修。", "TRN-02-006 未完成必修课会影响年度绩效流程提交。")),
            Section("外部培训", ("TRN-04-005 外部培训费用超过8000元需签署培训服务协议；入职未满一年参加的服务期为2年。", "TRN-04-006 外出参会产生的差旅费用按差旅制度报销，培训费用按培训预算审批。")),
            Section("晋升评审", ("TRN-05-002 晋升材料须包含近两期绩效、项目贡献和能力证明。", "TRN-05-008 晋升答辩未通过可在6个月后再次申请。")),
            Section("导师机制", ("TRN-06-004 导师每月至少完成一次成长面谈并记录行动项。", "TRN-06-007 导师评价不直接决定绩效，但作为晋升参考输入。")),
        ),
    ),
    EnterpriseDoc(
        "研发规范与发布流程",
        ".txt",
        "研发效能部",
        "规定需求评审、代码管理、测试准入、发布窗口和数据库变更。",
        (
            Section("需求评审", ("RD-01-002 需求进入开发前须明确验收标准、边界条件和回滚影响。", "RD-01-007 跨部门需求需指定一个业务负责人统一决策。")),
            Section("代码管理", ("RD-02-003 主干分支必须保持可构建，功能分支通过合并请求进入主干。", "RD-02-008 高风险模块至少两名维护人完成代码评审。")),
            Section("测试准入", ("RD-04-001 单元测试、接口测试和核心回归用例通过后方可提测。", "RD-04-006 缺陷等级S1/S2未关闭不得进入正式发布。")),
            Section("发布窗口与紧急发布", ("RD-05-001 生产发布窗口默认安排在周二、周四20:00-22:00。", "RD-05-006 紧急hotfix也必须提供验证记录和回滚方案，不得无回滚方案直接发布。")),
            Section("数据库变更", ("RD-06-004 生产数据库变更需DBA审核，涉及敏感数据时需信息安全复核，并由发布经理确认窗口。", "RD-06-009 数据库回滚脚本必须在预发环境演练通过。")),
        ),
    ),
    EnterpriseDoc(
        "行政后勤制度",
        ".docx",
        "行政部",
        "覆盖访客、办公区安全、食堂、会议室和固定资产管理。",
        (
            Section("访客管理", ("ADM-02-004 访客临时门禁最长有效期为8小时，超时需重新登记。", "ADM-02-007 访客进入研发区需由接待人全程陪同。")),
            Section("办公区安全", ("ADM-03-002 非工作时间进入办公区需登记原因。", "ADM-03-006 办公区拍摄涉及屏幕内容须获得部门负责人同意。")),
            Section("食堂管理", ("ADM-04-001 食堂补贴按月结算，不兑换现金。", "ADM-04-005 食品安全投诉须在当日反馈给供应商管理员。")),
            Section("会议室", ("ADM-05-003 超过30分钟未签到的会议室会自动释放。", "ADM-05-008 董事会会议室由行政专员统一排期。")),
            Section("固定资产", ("ADM-06-002 固定资产调拨需在系统登记接收人和位置。", "ADM-06-007 资产盘点差异须在5个工作日内完成复核。")),
        ),
    ),
    EnterpriseDoc(
        "财务付款与预算制度",
        ".csv",
        "财务部",
        "定义预算编制、预算执行、预算外付款、付款材料和费用归档。",
        (
            Section("预算编制", ("FIN-01-002 年度预算按部门、项目和费用类型三维度编制。", "FIN-01-006 预算调整需说明业务变化和资金来源。")),
            Section("预算执行", ("FIN-03-001 预算内付款需校验合同、验收和发票一致性。", "FIN-03-006 预算余额不足时系统不允许直接提交付款。")),
            Section("预算外付款", ("FIN-04-003 预算外付款超过10万元需CFO审批；超过或等于15万元需总经理会签。", "FIN-04-006 紧急预算外付款可先冻结额度后补齐审批。")),
            Section("付款材料", ("FIN-05-002 付款材料包括合同、验收单、发票、付款申请和收款账户证明。", "FIN-05-008 对公账户变更需供应商重新盖章确认。")),
            Section("费用归档", ("FIN-06-003 付款完成后电子档案保留不少于10年。", "FIN-06-007 审计抽样发现缺失材料时应在3个工作日内补齐。")),
        ),
    ),
    EnterpriseDoc(
        "合同审批制度",
        ".pdf",
        "法务部",
        "约束合同起草、评审时限、金额审批、用印和用印后变更。",
        (
            Section("合同起草", ("CON-01-001 合同应使用公司模板，偏离模板条款需标红说明。", "CON-01-006 涉及个人信息处理的合同须附数据处理协议。")),
            Section("合同评审时限", ("CON-03-002 标准合同法务初审时限为2个工作日。", "CON-03-006 非标合同初审时限为5个工作日，重大项目可另行约定。")),
            Section("合同金额审批", ("CON-04-003 标准合同金额80万元需业务负责人、法务和财务审批。", "CON-04-008 超过200万元合同需总经理审批。")),
            Section("用印管理", ("CON-06-002 用印前需确认审批链完整、附件齐备和合同版本一致。", "CON-06-006 空白合同、空白授权书不得用印。")),
            Section("用印后变更", ("CON-07-002 合同盖章后不得直接修改正文，确需修改应重新走变更审批。", "CON-07-006 已寄出的合同如需作废，应登记作废原因和快递追踪。")),
        ),
    ),
    EnterpriseDoc(
        "客户支持与售后制度",
        ".txt",
        "客户成功部",
        "覆盖故障响应、升级处理、客户回访、知识沉淀和服务报告。",
        (
            Section("故障响应", ("CS-02-001 S1客户故障需30分钟内给出处理方案，1小时内同步阶段性进展。", "CS-02-006 S2问题需在4小时内给出临时方案。")),
            Section("升级处理", ("CS-03-004 S1客户故障超过4小时未恢复，需升级业务负责人和CTO；涉及发布时研发走hotfix流程。", "CS-03-008 涉及合同赔付风险时应同步法务和客户经理。")),
            Section("客户回访", ("CS-04-002 重大故障关闭后3个工作日内完成客户回访。", "CS-04-006 客户满意度低于3分需制定改进计划。")),
            Section("知识沉淀", ("CS-05-001 高频问题需沉淀为知识库文章并关联产品版本。", "CS-05-006 临时绕行方案过期后需标记失效。")),
            Section("服务报告", ("CS-06-003 月度服务报告包括工单量、SLA达成率、故障复盘和待办风险。", "CS-06-008 战略客户服务报告需客户成功负责人复核。")),
        ),
    ),
)


def _clean_output() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for item in OUT.iterdir():
        if item.is_file():
            item.unlink()
        elif item.is_dir():
            shutil.rmtree(item)


def _expansion(doc: EnterpriseDoc, section: Section, idx: int) -> list[str]:
    return [
        f"责任边界：{section.title}由{doc.owner}主责，涉及跨部门事项时同步财务、法务、IT或信息安全。",
        f"审批证据：第{idx}类事项至少保留审批单、系统截图、沟通纪要和最终处理结果。",
        f"例外规则：当{section.title}与其他制度适用边界不清时，先冻结流程状态，由归口部门出具书面判断。",
        f"审计要求：季度抽查会比对条款编号、审批人、时间戳和业务单据，异常项进入整改台账。",
        f"评测锚点：该章节用于测试文档命中、章节命中、条款编号召回和跨文档引用能力。",
    ]


def _page_title(doc: EnterpriseDoc, page_no: int) -> str:
    section = doc.sections[(page_no - 1) % len(doc.sections)]
    return f"第{page_no:02d}页 {doc.name} - {section.title}"


def _page_paragraphs(doc: EnterpriseDoc, page_no: int) -> list[str]:
    section = doc.sections[(page_no - 1) % len(doc.sections)]
    cycle = (page_no - 1) // len(doc.sections) + 1
    lines = [
        f"页码：第{page_no:02d}页",
        f"章节：{section.title}",
        f"版本场景：第{cycle}轮业务复盘，覆盖制度解释、执行证据、边界条件和异常处理。",
        f"业务背景：{doc.owner}在处理{section.title}事项时，需要同时考虑申请人、审批人、执行人和审计人的职责分离。",
    ]
    lines.extend(section.clauses)
    lines.extend(_expansion(doc, section, page_no))
    scenario_templates = [
        "申请入口：员工或业务负责人在系统中提交事项，并选择费用、权限、合同、供应商或客户影响范围。",
        "审批规则：系统根据金额、数据等级、故障等级、合同类型或员工层级自动识别审批链。",
        "执行动作：责任团队在规定时限内完成处理，并把执行截图、邮件确认和系统日志关联到单据。",
        "例外判断：遇到客户紧急故障、预算冻结、供应商风险或安全事件时，先采取临时控制，再补齐审批。",
        "审计抽查：内控团队按月抽样检查条款编号、审批时点、处理结果和证据完整性。",
        "跨文档引用：涉及付款时参考财务制度，涉及合同条款时参考合同审批制度，涉及生产系统时参考IT和研发制度。",
        "失败风险：若只命中文档标题但未召回本页条款，评测应归类为章节未命中或答案要点缺失。",
        "整改要求：发现执行偏差后，责任部门应在5个工作日内提交原因、影响范围和防复发措施。",
    ]
    for idx, template in enumerate(scenario_templates, start=1):
        lines.append(f"{doc.name}-P{page_no:02d}-S{idx:02d} {template}")
    return lines


def _text_lines(doc: EnterpriseDoc) -> list[str]:
    lines = [f"# {doc.name}", f"归口部门：{doc.owner}", f"适用范围：{doc.scope}", ""]
    for page_no in range(1, PAGE_COUNT + 1):
        lines.extend([f"## {_page_title(doc, page_no)}", ""])
        lines.extend(_page_paragraphs(doc, page_no))
        lines.append("")
    return lines


def _write_txt(doc: EnterpriseDoc) -> None:
    (OUT / f"{doc.name}{doc.suffix}").write_text("\n".join(_text_lines(doc)), encoding="utf-8")


def _write_csv(doc: EnterpriseDoc) -> None:
    with (OUT / f"{doc.name}{doc.suffix}").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["文档", "页码", "章节", "条款编号", "条款内容", "责任部门", "证据材料", "例外处理"])
        for page_no in range(1, PAGE_COUNT + 1):
            section = doc.sections[(page_no - 1) % len(doc.sections)]
            for clause in section.clauses:
                clause_id, text = clause.split(" ", 1)
                writer.writerow([doc.name, f"第{page_no:02d}页", section.title, clause_id, text, doc.owner, "审批单/系统截图/会议纪要", "升级归口部门复核"])
            for note in _page_paragraphs(doc, page_no):
                writer.writerow([doc.name, f"第{page_no:02d}页", section.title, "", note, doc.owner, "执行记录", "按更严格制度处理"])


def _write_docx(doc: EnterpriseDoc) -> None:
    document = DocxDocument()
    document.add_heading(doc.name, 0)
    document.add_paragraph(f"归口部门：{doc.owner}")
    document.add_paragraph(f"适用范围：{doc.scope}")
    document.add_paragraph(f"页数控制：本文档按 {PAGE_COUNT} 页等价内容生成，每页包含独立业务场景和可检索条款。")
    for page_no in range(1, PAGE_COUNT + 1):
        section = doc.sections[(page_no - 1) % len(doc.sections)]
        document.add_heading(_page_title(doc, page_no), level=1)
        for clause in section.clauses:
            document.add_paragraph(clause, style="List Bullet")
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for cell, label in zip(table.rows[0].cells, ["控制点", "责任人", "证据", "风险提示"]):
            cell.text = label
        for note in _page_paragraphs(doc, page_no):
            row = table.add_row().cells
            row[0].text = section.title
            row[1].text = doc.owner
            row[2].text = "审批记录、系统日志、验收材料"
            row[3].text = note
        document.add_paragraph(f"跨文档引用：若{section.title}涉及采购、付款、合同或生产发布，应同步查看对应制度的审批边界。")
        if page_no < PAGE_COUNT:
            document.add_page_break()
    for section in document.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
    document.save(OUT / f"{doc.name}{doc.suffix}")


def _write_xlsx(doc: EnterpriseDoc) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "评测明细"
    summary.append(["文档", "页码", "章节", "条款编号", "条款内容", "责任部门", "审批/留痕", "评测关键词"])
    for page_no in range(1, PAGE_COUNT + 1):
        section = doc.sections[(page_no - 1) % len(doc.sections)]
        for clause in section.clauses:
            clause_id, text = clause.split(" ", 1)
            summary.append([doc.name, f"第{page_no:02d}页", section.title, clause_id, text, doc.owner, "系统审批、验收材料、审计日志", "时限/金额/角色/例外"])
        for note in _page_paragraphs(doc, page_no):
            summary.append([doc.name, f"第{page_no:02d}页", section.title, "补充说明", note, doc.owner, "复核记录", section.title])
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for col in ("A", "B", "C", "D", "E", "F", "G", "H"):
        summary.column_dimensions[col].width = 22 if col != "E" else 84
    for row in summary.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for page_no in range(1, PAGE_COUNT + 1):
        section = doc.sections[(page_no - 1) % len(doc.sections)]
        sheet = workbook.create_sheet(f"P{page_no:02d}-{section.title}"[:31])
        sheet.append(["页码", "章节", "条款编号", "条款内容", "责任部门", "审批/留痕", "评测关键词"])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(wrap_text=True)
        for clause in section.clauses:
            clause_id, text = clause.split(" ", 1)
            sheet.append([f"第{page_no:02d}页", section.title, clause_id, text, doc.owner, "系统审批、验收材料、审计日志", "时限/金额/角色/例外"])
        for note in _page_paragraphs(doc, page_no):
            sheet.append([f"第{page_no:02d}页", section.title, "补充说明", note, doc.owner, "复核记录", section.title])
        for col in ("A", "B", "C", "D", "E", "F", "G"):
            sheet.column_dimensions[col].width = 24 if col != "D" else 84
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(OUT / f"{doc.name}{doc.suffix}")


def _font_path() -> Path:
    candidates = [
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
    ]
    for path in candidates:
        if path.exists():
            return path
    raise RuntimeError("未找到可用于生成中文 PDF 的字体，请安装微软雅黑、黑体或 Noto Sans CJK。")


def _write_pdf(doc: EnterpriseDoc) -> None:
    pdfmetrics.registerFont(TTFont("CJKFont", str(_font_path())))
    path = OUT / f"{doc.name}{doc.suffix}"
    pdf = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4

    for page_no in range(1, PAGE_COUNT + 1):
        section = doc.sections[(page_no - 1) % len(doc.sections)]
        y = height - 18 * mm
        pdf.setFont("CJKFont", 15)
        pdf.drawString(18 * mm, y, doc.name)
        y -= 8 * mm
        pdf.setFont("CJKFont", 10)
        pdf.drawString(18 * mm, y, _page_title(doc, page_no))
        y -= 6 * mm
        pdf.drawString(18 * mm, y, f"归口部门：{doc.owner}")
        y -= 6 * mm
        pdf.drawString(18 * mm, y, f"适用范围：{doc.scope}")
        y -= 8 * mm
        pdf.line(18 * mm, y, width - 18 * mm, y)
        y -= 7 * mm

        for line in _page_paragraphs(doc, page_no):
            prefix = "条款：" if any(line.startswith(clause.split(" ", 1)[0]) for clause in section.clauses) else "说明："
            for wrapped in _wrap_pdf_line(f"{prefix}{line}", max_chars=43):
                pdf.drawString(18 * mm, y, wrapped)
                y -= 5.6 * mm
                if y < 24 * mm:
                    break
            if y < 24 * mm:
                break

        pdf.setFont("CJKFont", 8)
        pdf.drawRightString(width - 18 * mm, 12 * mm, f"{doc.name} 第 {page_no} / {PAGE_COUNT} 页")
        pdf.showPage()
    pdf.save()


def _wrap_pdf_line(text: str, max_chars: int) -> list[str]:
    chunks: list[str] = []
    current = ""
    for char in text:
        current += char
        if len(current) >= max_chars:
            chunks.append(current)
            current = ""
    if current:
        chunks.append(current)
    return chunks


WRITERS = {
    ".txt": _write_txt,
    ".csv": _write_csv,
    ".docx": _write_docx,
    ".xlsx": _write_xlsx,
    ".pdf": _write_pdf,
}


def main() -> None:
    _clean_output()
    for doc in DOCS:
        WRITERS[doc.suffix](doc)
    print(f"Generated {len(DOCS)} files in {OUT}")


if __name__ == "__main__":
    main()
